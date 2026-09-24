import csv
import re
from decimal import Decimal
from io import StringIO
from pathlib import Path
from datetime import datetime, timezone

from openpyxl import load_workbook
from pyxlsb import open_workbook as open_xlsb_workbook
import xlrd
from sqlalchemy import func, or_

from main.common.errors import AppError
from main.database.extensions import db
from main.modules.code_descriptions.service import CodeDescriptionService, normalize_code, reason_lookup_codes
from main.modules.auth.models import User
from main.modules.denials.models import AssignmentRule, Denial, DenialAction, DenialCategory
from main.modules.denials.utils import normalize_header, parse_date, parse_decimal, parse_service_dates, validate_headers


class DenialService:
    @staticmethod
    def list_denials(user, status="pending"):
        query = Denial.query.filter_by(organization_id=user.organization_id, appeal_status=status)
        if user.role != "admin":
            query = query.filter(or_(Denial.assigned_to_id == user.id, Denial.assigned_to_id.is_(None)))
        denials = query.order_by(Denial.created_at.desc()).all()
        return DenialService._with_reference_descriptions(denials)

    @staticmethod
    def list_pending(user):
        return DenialService.list_denials(user, "pending")

    @staticmethod
    def list_in_process(user):
        return DenialService.list_denials(user, "in_process")

    @staticmethod
    def list_completed(user):
        return DenialService.list_denials(user, "completed")

    @staticmethod
    def _with_reference_descriptions(denials):
        reason_codes = []
        remark_codes = []
        for denial in denials:
            reason_codes.extend(DenialService._split_codes(denial.reason))
            remark_codes.extend(DenialService._split_codes(denial.remark))

        reason_descriptions = CodeDescriptionService.description_map("reason", reason_codes)
        remark_descriptions = CodeDescriptionService.description_map("remark", remark_codes)
        category_codes = {
            DenialService._denial_category_code(code)
            for code in reason_codes
            if DenialService._denial_category_code(code)
        }
        category_map = DenialService._denial_category_map(category_codes)

        enriched_denials = []
        for denial in denials:
            item = denial.to_dict()
            item["reason_desc"] = DenialService._description_text(
                "reason",
                denial.reason,
                reason_descriptions,
            )
            item["remark_desc"] = DenialService._description_text(
                "remark",
                denial.remark,
                remark_descriptions,
            )
            item.update(DenialService._category_payload(denial.reason, category_map))
            enriched_denials.append(item)
        return enriched_denials

    @staticmethod
    def _denial_category_code(code):
        text = str(code or "").strip().upper()
        compact = text.replace(" ", "")
        match = re.match(r"^(?:CO|PR|PI|OA)-?(\d+)$", compact)
        if match:
            return match.group(1)
        return compact or None

    @staticmethod
    def _denial_category_map(codes):
        if not codes:
            return {}
        rows = DenialCategory.query.filter(DenialCategory.code.in_(codes)).all()
        return {row.code: row.category for row in rows}

    @staticmethod
    def _category_payload(value, category_map):
        details = []
        categories = []
        for code in DenialService._split_codes(value):
            lookup_code = DenialService._denial_category_code(code)
            category = category_map.get(lookup_code) if lookup_code else None
            resolved_category = category or "Uncategorized"
            details.append({
                "code": code,
                "lookup_code": lookup_code,
                "category": resolved_category,
            })
            categories.append(resolved_category)

        if not details:
            categories = ["Uncategorized"]

        unique_categories = sorted(set(categories), key=lambda item: item.lower())
        return {
            "denial_category": ", ".join(unique_categories),
            "denial_categories": unique_categories,
            "denial_category_details": details,
        }

    @staticmethod
    def _description_text(code_type, value, description_map):
        entries = []
        for code in DenialService._split_codes(value):
            lookup_codes = reason_lookup_codes(code) if code_type == "reason" else [normalize_code(code_type, code)]
            description = next((description_map.get(item) for item in lookup_codes if description_map.get(item)), None)
            if description:
                entries.append(f"{code}: {description}")
        return "\n".join(entries)

    @staticmethod
    def preview_excel(admin_user, file_storage):
        denials = DenialService._parse_excel(admin_user, file_storage)
        analysis = DenialService._analyze_denials(admin_user, denials)
        return {
            "row_count": len(denials),
            "exact_duplicates": analysis["exact_duplicates"],
            "second_time_denials": analysis["second_time_denials"],
            "fresh_denials": analysis["fresh_denials"],
        }

    @staticmethod
    def upload_excel(admin_user, file_storage):
        denials = DenialService._parse_excel(admin_user, file_storage)
        analysis = DenialService._analyze_denials(admin_user, denials)
        uploadable_denials = [item["denial"] for item in analysis["classified_rows"] if item["classification"] != "exact"]
        DenialService._apply_assignment_rules(admin_user.organization_id, uploadable_denials)
        if uploadable_denials:
            db.session.add_all(uploadable_denials)
        db.session.commit()
        return {
            "uploaded_count": len(uploadable_denials),
            "discarded_duplicates": analysis["exact_duplicates"],
            "second_time_denials": analysis["second_time_denials"],
            "fresh_denials": analysis["fresh_denials"],
        }

    @staticmethod
    def _parse_excel(admin_user, file_storage):
        if not file_storage:
            raise AppError("Upload a spreadsheet file.", 422)

        rows = DenialService._read_rows(file_storage)
        if not rows:
            raise AppError("The uploaded file is empty.", 422)

        headers = [str(header or "").strip() for header in rows[0]]
        validate_headers(headers)
        header_index = {normalize_header(header): index for index, header in enumerate(headers)}

        denials = []
        for row in rows[1:]:
            if not any(row):
                continue
            value = lambda column: row[header_index[normalize_header(column)]] if header_index[normalize_header(column)] < len(row) else None
            patient = str(value("Patient") or "").strip()
            claim_id = str(value("Claim ID") or "").strip()
            if not patient or not claim_id:
                continue
            service_date, service_date_display = parse_service_dates(value("Service Date"))
            denials.append(
                Denial(
                    organization_id=admin_user.organization_id if admin_user else 0,
                    patient=patient,
                    claim_id=claim_id,
                    service_date=service_date,
                    service_date_display=service_date_display,
                    billed_amt=parse_decimal(value("Billed Amt")),
                    payment_amt=parse_decimal(value("Payment Amt")),
                    payer=str(value("Payer") or "").strip() or None,
                    provider=str(value("Provider") or "").strip() or None,
                    practice=str(value("Practice") or "").strip() or None,
                    procedure=str(value("Procedure") or "").strip() or None,
                    modifier=str(value("Modifier") or "").strip() or None,
                    reason=str(value("Reason") or "").strip() or None,
                    remark=str(value("Remark") or "").strip() or None,
                    assigned_to_id=None,
                    eob_date=parse_date(value("EOB Date")),
                    notes=str(value("Notes") or "").strip() or None,
                    date_added=parse_date(value("Date Added")),
                )
            )

        if not denials:
            raise AppError("No valid denial rows were found.", 422)
        return denials

    @staticmethod
    def _analyze_denials(admin_user, denials):
        existing_denials = Denial.query.filter_by(organization_id=admin_user.organization_id).all()
        exact_seen = {DenialService._exact_signature(denial) for denial in existing_denials}
        base_seen = {DenialService._base_signature(denial) for denial in existing_denials}

        classified_rows = []
        counts = {"exact_duplicates": 0, "second_time_denials": 0, "fresh_denials": 0}

        for denial in denials:
            exact_signature = DenialService._exact_signature(denial)
            base_signature = DenialService._base_signature(denial)
            if exact_signature in exact_seen:
                classification = "exact"
                counts["exact_duplicates"] += 1
            elif base_signature in base_seen:
                classification = "second_time"
                counts["second_time_denials"] += 1
            else:
                classification = "fresh"
                counts["fresh_denials"] += 1

            classified_rows.append({"denial": denial, "classification": classification})
            exact_seen.add(exact_signature)
            base_seen.add(base_signature)

        return {**counts, "classified_rows": classified_rows}

    @staticmethod
    def _normalize_signature_value(value):
        if value is None:
            return ""
        if isinstance(value, Decimal):
            return str(value.quantize(Decimal("0.01")))
        return str(value).strip().lower()

    @staticmethod
    def _exact_signature(denial):
        return tuple(
            DenialService._normalize_signature_value(value)
            for value in (
                denial.patient,
                denial.claim_id,
                denial.service_date,
                denial.service_date_display,
                denial.billed_amt,
                denial.payment_amt,
                denial.payer,
                denial.provider,
                denial.practice,
                denial.procedure,
                denial.modifier,
                denial.reason,
                denial.remark,
                denial.eob_date,
                denial.notes,
                denial.date_added,
            )
        )

    @staticmethod
    def _base_signature(denial):
        return tuple(
            DenialService._normalize_signature_value(value)
            for value in (
                denial.patient,
                denial.claim_id,
                denial.service_date,
                denial.service_date_display,
                denial.billed_amt,
                denial.payment_amt,
                denial.payer,
                denial.provider,
                denial.practice,
                denial.procedure,
                denial.modifier,
                denial.eob_date,
                denial.date_added,
            )
        )

    @staticmethod
    def _read_rows(file_storage):
        extension = Path(file_storage.filename or "").suffix.lower()
        file_storage.stream.seek(0)

        if extension in (".xlsx", ".xlsm", ".xltx", ".xltm"):
            workbook = load_workbook(file_storage.stream, read_only=True, data_only=True)
            sheet = workbook.active
            return list(sheet.iter_rows(values_only=True))

        if extension == ".xls":
            workbook = xlrd.open_workbook(file_contents=file_storage.read())
            sheet = workbook.sheet_by_index(0)
            return [sheet.row_values(index) for index in range(sheet.nrows)]

        if extension == ".xlsb":
            rows = []
            with open_xlsb_workbook(file_storage.stream) as workbook:
                with workbook.get_sheet(1) as sheet:
                    for row in sheet.rows():
                        rows.append([cell.v for cell in row])
            return rows

        if extension in (".csv", ".txt", ".tsv"):
            delimiter = "\t" if extension == ".tsv" else ","
            content = file_storage.read().decode("utf-8-sig")
            return list(csv.reader(StringIO(content), delimiter=delimiter))

        raise AppError("Supported formats are .csv, .tsv, .txt, .xlsx, .xlsm, .xltx, .xltm, .xls, and .xlsb.", 422)

    @staticmethod
    def assign_denial(admin_user, denial_id, user_id):
        denial = Denial.query.filter_by(id=denial_id, organization_id=admin_user.organization_id).first()
        if not denial:
            raise AppError("Denial not found.", 404)
        assignee = User.query.filter_by(id=user_id, organization_id=admin_user.organization_id, is_active=True).first()
        if not assignee:
            raise AppError("Assignee not found.", 404)
        denial.assigned_to_id = assignee.id
        db.session.commit()
        return denial.to_dict()

    @staticmethod
    def list_actions():
        return [
            action.to_dict()
            for action in DenialAction.query.filter_by(is_active=True).order_by(DenialAction.id.asc()).all()
        ]

    @staticmethod
    def create_action(admin_user, data):
        name = str(data.get("name") or "").strip()
        if not name:
            raise AppError("Enter an action name.", 422)
        existing = DenialAction.query.filter(func.lower(DenialAction.name) == name.lower()).first()
        if existing:
            if not existing.is_active:
                existing.is_active = True
                db.session.commit()
            return existing.to_dict()

        action = DenialAction(name=name, is_active=True)
        db.session.add(action)
        db.session.commit()
        return action.to_dict()

    @staticmethod
    def assign_action(user, denial_id, data):
        denial = DenialService._denial_for_user(user, denial_id)
        action = DenialAction.query.filter_by(id=data["action_id"], is_active=True).first()
        if not action:
            raise AppError("Select a valid action.", 422)

        other_description = str(data.get("other_description") or "").strip()
        if action.name.lower() == "other" and not other_description:
            raise AppError("Enter a description for Other action.", 422)

        denial.assigned_action_id = action.id
        denial.action_other_description = other_description or None
        denial.action_notes = str(data.get("notes") or "").strip() or None
        denial.action_assigned_by_id = user.id
        denial.action_completed_by_id = None
        denial.action_assigned_at = datetime.now(timezone.utc)
        denial.action_completed_at = None
        denial.appeal_status = "in_process"
        denial.appeal_stage = 1
        denial.appeal_started_at = None
        denial.appeal_completed_at = None
        denial.appeal_started_by_id = None
        denial.appeal_completed_by_id = None
        db.session.commit()
        return denial.to_dict()

    @staticmethod
    def remove_action(user, denial_id):
        denial = DenialService._denial_for_user(user, denial_id)
        if not denial.assigned_action_id:
            raise AppError("No action is assigned to this denial.", 422)
        DenialService._clear_action(denial)
        denial.appeal_status = "pending"
        db.session.commit()
        return denial.to_dict()

    @staticmethod
    def complete_action(user, denial_id):
        denial = DenialService._denial_for_user(user, denial_id)
        if not denial.assigned_action_id:
            raise AppError("No action is assigned to this denial.", 422)
        denial.appeal_status = "completed"
        denial.action_completed_by_id = user.id
        denial.action_completed_at = datetime.now(timezone.utc)
        db.session.commit()
        return denial.to_dict()

    @staticmethod
    def _clear_action(denial):
        denial.assigned_action_id = None
        denial.action_other_description = None
        denial.action_notes = None
        denial.action_assigned_by_id = None
        denial.action_completed_by_id = None
        denial.action_assigned_at = None
        denial.action_completed_at = None

    @staticmethod
    def start_appeal(user, denial_id):
        denial = DenialService._denial_for_user(user, denial_id)
        DenialService._clear_action(denial)
        denial.appeal_status = "in_process"
        denial.appeal_stage = max(1, denial.appeal_stage or 1)
        if not denial.appeal_started_at:
            denial.appeal_started_at = datetime.now(timezone.utc)
        denial.appeal_completed_at = None
        denial.appeal_started_by_id = user.id
        denial.appeal_completed_by_id = None
        db.session.commit()
        return denial.to_dict()

    @staticmethod
    def update_appeal_stage(user, denial_id, stage):
        denial = DenialService._denial_for_user(user, denial_id)
        try:
            stage_number = int(stage)
        except (TypeError, ValueError) as exc:
            raise AppError("Select a valid appeal stage.", 422) from exc
        if stage_number < 1 or stage_number > 5:
            raise AppError("Appeal stage must be between 1 and 5.", 422)
        denial.appeal_stage = stage_number
        db.session.commit()
        return denial.to_dict()

    @staticmethod
    def remove_appeal(user, denial_id):
        denial = DenialService._denial_for_user(user, denial_id)
        if denial.assigned_action_id:
            return DenialService.remove_action(user, denial_id)
        denial.appeal_status = "pending"
        denial.appeal_stage = 1
        denial.appeal_started_at = None
        denial.appeal_completed_at = None
        denial.appeal_completed_by_id = None
        db.session.commit()
        return denial.to_dict()

    @staticmethod
    def complete_appeal(user, denial_id):
        denial = DenialService._denial_for_user(user, denial_id)
        DenialService._clear_action(denial)
        denial.appeal_status = "completed"
        denial.appeal_stage = 5
        denial.appeal_completed_at = datetime.now(timezone.utc)
        denial.appeal_completed_by_id = user.id
        db.session.commit()
        return denial.to_dict()

    @staticmethod
    def _denial_for_user(user, denial_id):
        query = Denial.query.filter_by(id=denial_id, organization_id=user.organization_id)
        if user.role != "admin":
            query = query.filter(or_(Denial.assigned_to_id == user.id, Denial.assigned_to_id.is_(None)))
        denial = query.first()
        if not denial:
            raise AppError("Denial not found.", 404)
        return denial

    @staticmethod
    def list_assignment_rules(user):
        query = AssignmentRule.query.filter_by(organization_id=user.organization_id)
        if user.role != "admin":
            query = query.filter_by(assigned_to_id=user.id)
        return [rule.to_dict() for rule in query.order_by(AssignmentRule.created_at.desc()).all()]

    @staticmethod
    def create_assignment_rule(admin_user, data):
        assignee = User.query.filter_by(
            id=data["assigned_to_id"],
            organization_id=admin_user.organization_id,
            is_active=True,
        ).first()
        if not assignee:
            raise AppError("Rule assignee not found.", 404)

        rule = AssignmentRule(
            organization_id=admin_user.organization_id,
            assigned_to_id=assignee.id,
            created_by_id=admin_user.id,
            criteria=data["criteria"],
            effective_from=DenialService._effective_datetime(data.get("effective_from")),
        )
        db.session.add(rule)
        db.session.flush()

        matching_denials = Denial.query.filter_by(organization_id=admin_user.organization_id).all()
        assigned_count = 0
        for denial in matching_denials:
            if DenialService._rule_is_effective(rule) and DenialService._rule_matches_denial(rule, denial):
                denial.assigned_to_id = assignee.id
                assigned_count += 1

        db.session.commit()
        return {"rule": rule.to_dict(), "assigned_count": assigned_count}

    @staticmethod
    def deactivate_assignment_rule(admin_user, rule_id):
        rule = AssignmentRule.query.filter_by(
            id=rule_id,
            organization_id=admin_user.organization_id,
        ).first()
        if not rule:
            raise AppError("Assignment rule not found.", 404)
        rule.is_active = False
        rule.deactivated_at = datetime.now(timezone.utc)
        db.session.commit()
        return rule.to_dict()

    @staticmethod
    def update_assignment_rule(admin_user, rule_id, data):
        rule = AssignmentRule.query.filter_by(
            id=rule_id,
            organization_id=admin_user.organization_id,
        ).first()
        if not rule:
            raise AppError("Assignment rule not found.", 404)
        rule.criteria = data["criteria"]
        rule.effective_from = DenialService._effective_datetime(data.get("effective_from"))
        if DenialService._rule_is_effective(rule):
            matching_denials = Denial.query.filter_by(organization_id=admin_user.organization_id).all()
            for denial in matching_denials:
                if DenialService._rule_matches_denial(rule, denial):
                    denial.assigned_to_id = rule.assigned_to_id
        db.session.commit()
        return rule.to_dict()

    @staticmethod
    def _apply_assignment_rules(organization_id, denials):
        rules = AssignmentRule.query.filter_by(organization_id=organization_id, is_active=True).order_by(
            AssignmentRule.created_at.asc()
        ).all()
        for denial in denials:
            for rule in rules:
                if DenialService._rule_is_effective(rule) and DenialService._rule_matches_denial(rule, denial):
                    denial.assigned_to_id = rule.assigned_to_id
                    break

    @staticmethod
    def _effective_datetime(value):
        if not value:
            return datetime.now(timezone.utc)
        return datetime.combine(value, datetime.min.time(), tzinfo=timezone.utc)

    @staticmethod
    def _rule_is_effective(rule):
        return not rule.effective_from or rule.effective_from <= datetime.now(timezone.utc)

    @staticmethod
    def _split_codes(value):
        if not value:
            return []
        normalized = str(value).replace(";", ",").replace("|", ",")
        return [item.strip().lower() for item in normalized.split(",") if item.strip()]

    @staticmethod
    def _matches_values(denial_values, rule_values):
        if not rule_values:
            return True
        normalized_denial_values = {str(value).strip().lower() for value in denial_values if value}
        normalized_rule_values = {str(value).strip().lower() for value in rule_values if value}
        return bool(normalized_denial_values.intersection(normalized_rule_values))

    @staticmethod
    def _rule_matches_denial(rule, denial):
        criteria = rule.criteria or {}
        checks = (
            ([denial.payer], criteria.get("payer", [])),
            ([denial.practice], criteria.get("practice", [])),
            (DenialService._split_codes(denial.procedure), criteria.get("procedure", [])),
            (DenialService._split_codes(denial.modifier), criteria.get("modifier", [])),
            (DenialService._split_codes(denial.reason), criteria.get("reason", [])),
            (DenialService._split_codes(denial.remark), criteria.get("remark", [])),
        )
        return all(DenialService._matches_values(denial_values, rule_values) for denial_values, rule_values in checks)
